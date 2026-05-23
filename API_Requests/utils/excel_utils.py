import os
from openpyxl import load_workbook, Workbook


class ExcelHandler:
    """
    读写excel文件
    """

    def __init__(self, excel_file):
        self.excel_file = excel_file

    def read_excel(self, sheet_name=None):
        """
        读excel文件
        """

        if not os.path.exists(self.excel_file):
            raise FileNotFoundError(f"文件{self.excel_file}不存在")

        wb = load_workbook(self.excel_file)
        if sheet_name:
            if sheet_name not in wb.sheetnames:
                raise ValueError(f"工作表{sheet_name}不存在")
            ws = wb[sheet_name]
        else:
            ws = wb.active

        result = []
        for row in range(2, ws.max_row + 1):
            data = []
            for col in range(1, ws.max_column + 1):
                data.append(ws.cell(row=row, column=col).value)
            result.append(data)

        wb.close()
        return result

    def write_excel(self, sheet_name=None, header=None, data_list: list = None):
        """
        写excel文件
        """

        if not os.path.exists(self.excel_file):
            wb = Workbook()
        else:
            wb = load_workbook(self.excel_file)

        if sheet_name:
            if sheet_name not in wb.sheetnames:
                wb.create_sheet(sheet_name)
            ws = wb[sheet_name]
        else:
            ws = wb.active

        if header:
            ws.append(header)

        if data_list and isinstance(data_list, list):
            if data_list and isinstance(data_list[0], list):
                for row in data_list:
                    ws.append(row)
            else:
                ws.append(data_list)

        wb.save(self.excel_file)
        wb.close()

